import json
import logging
import os
import openpyxl
from sqlalchemy import select, func
from app.db.models import LocationCoordinate
import app.db.session as db_session

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("seed_locations")

# File paths
WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
AIRPORTS_FILE = os.path.join(WORKSPACE_DIR, "..", "global_cargo_airports.xlsx")
if not os.path.exists(AIRPORTS_FILE):
    AIRPORTS_FILE = "/Users/nilesh/Downloads/tracking/global_cargo_airports.xlsx"

SEAPORTS_FILE = os.path.join(WORKSPACE_DIR, "..", "world_seaports_with_coordinates.xlsx")
if not os.path.exists(SEAPORTS_FILE):
    SEAPORTS_FILE = "/Users/nilesh/Downloads/tracking/world_seaports_with_coordinates.xlsx"

# Modern container transshipment ports commonly encountered in live carrier feeds
SUPPLEMENTARY_PORTS = [
    {
        "location_type": "SEAPORT",
        "code": "INMUN",
        "display_code": "IN MUN",
        "icao_code": None,
        "name": "Mundra Port",
        "city": "Mundra",
        "country": "India",
        "latitude": 22.8396,
        "longitude": 69.7042,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Adani Mundra Port & SEZ"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "QAHMD",
        "display_code": "QA HMD",
        "icao_code": None,
        "name": "Hamad Port",
        "city": "Hamad",
        "country": "Qatar",
        "latitude": 25.0131,
        "longitude": 51.6114,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Doha Hamad Mega Port"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "ITGIT",
        "display_code": "IT GIT",
        "icao_code": None,
        "name": "Port of Gioia Tauro",
        "city": "Gioia Tauro",
        "country": "Italy",
        "latitude": 38.4333,
        "longitude": 15.9000,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Mediterranean Hub"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "INGGV",
        "display_code": "IN GGV",
        "icao_code": None,
        "name": "Gangavaram Port",
        "city": "Gangavaram",
        "country": "India",
        "latitude": 17.6186,
        "longitude": 83.2383,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Vizag Gangavaram Deepwater Port"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "INTRV",
        "display_code": "IN TRV",
        "icao_code": None,
        "name": "Vizhinjam International Seaport",
        "city": "Vizhinjam",
        "country": "India",
        "latitude": 8.3756,
        "longitude": 76.9917,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Adani Vizhinjam Hub"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "INKAT",
        "display_code": "IN KAT",
        "icao_code": None,
        "name": "Kattupalli Port",
        "city": "Kattupalli",
        "country": "India",
        "latitude": 13.3100,
        "longitude": 80.3400,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Adani Kattupalli Port"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "INENR",
        "display_code": "IN ENR",
        "icao_code": None,
        "name": "Kamarajar Port (Ennore)",
        "city": "Ennore",
        "country": "India",
        "latitude": 13.2608,
        "longitude": 80.3340,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Kamarajar Port Ennore"}),
    },
    {
        "location_type": "SEAPORT",
        "code": "RULED",
        "display_code": "RU LED",
        "icao_code": None,
        "name": "Saint Petersburg Port",
        "city": "Saint Petersburg",
        "country": "Russia",
        "latitude": 59.88333,
        "longitude": 30.25,
        "source": "supplementary_carrier_master",
        "metadata_json": json.dumps({"notes": "Sea Port of Saint Petersburg"}),
    },
]


def load_airports(filepath: str):
    logger.info(f"Loading airports from {filepath}...")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb["All Cargo Airports"]
    records = []
    seen = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 8:
            continue
        country, name, lat, lon, airport_code, iata_code, icao_code, city = row[0:8]
        cargo_status = row[8] if len(row) > 8 else None
        category = row[9] if len(row) > 9 else None
        airport_type = row[10] if len(row) > 10 else None
        scheduled = row[11] if len(row) > 11 else None
        continent = row[12] if len(row) > 12 else None

        if lat is None or lon is None:
            continue

        try:
            lat_f = float(lat)
            lon_f = float(lon)
        except (ValueError, TypeError):
            continue

        primary_code = (str(iata_code).strip().upper() if iata_code else None) or (str(airport_code).strip().upper() if airport_code else None) or (str(icao_code).strip().upper() if icao_code else None)
        if not primary_code or primary_code in ("NULL", "NONE", "-"):
            continue

        icao_str = str(icao_code).strip().upper() if icao_code and str(icao_code).strip().upper() not in ("NULL", "NONE", "-") else None

        key = ("AIRPORT", primary_code)
        if key in seen:
            continue
        seen.add(key)

        metadata = {
            "cargo_status": cargo_status,
            "category": category,
            "airport_type": airport_type,
            "scheduled": scheduled,
            "continent": continent,
        }

        records.append({
            "location_type": "AIRPORT",
            "code": primary_code,
            "display_code": primary_code,
            "icao_code": icao_str,
            "name": str(name).strip() if name else f"Airport {primary_code}",
            "city": str(city).strip() if city else None,
            "country": str(country).strip() if country else None,
            "latitude": lat_f,
            "longitude": lon_f,
            "source": "global_cargo_airports.xlsx",
            "metadata_json": json.dumps(metadata),
        })

    logger.info(f"Parsed {len(records)} unique airports.")
    return records


def load_seaports(filepath: str):
    logger.info(f"Loading seaports from {filepath}...")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb["World Seaports"]
    records = []
    seen = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 5:
            continue
        country, port, code, lat, lon = row[0:5]

        if lat is None or lon is None:
            continue

        try:
            lat_f = float(lat)
            lon_f = float(lon)
        except (ValueError, TypeError):
            continue

        display_code = str(code).strip() if code else ""
        norm_code = display_code.replace(" ", "").upper()
        if not norm_code or norm_code in ("NULL", "NONE", "-"):
            continue

        key = ("SEAPORT", norm_code)
        if key in seen:
            continue
        seen.add(key)

        records.append({
            "location_type": "SEAPORT",
            "code": norm_code,
            "display_code": display_code,
            "icao_code": None,
            "name": str(port).strip() if port else f"Port {norm_code}",
            "city": str(port).strip() if port else None,
            "country": str(country).strip() if country else None,
            "latitude": lat_f,
            "longitude": lon_f,
            "source": "world_seaports_with_coordinates.xlsx",
            "metadata_json": json.dumps({"notes": "World Seaports Master"}),
        })

    logger.info(f"Parsed {len(records)} unique seaports.")
    return records


def seed_locations_if_needed(force: bool = False):
    if db_session.SessionLocal is None:
        db_session.init_db_engine()
    db = db_session.SessionLocal()
    try:
        count = db.execute(select(func.count(LocationCoordinate.id))).scalar() or 0
        if count > 0 and not force:
            logger.debug(f"location_coordinates table already has {count} entries. Skipping seed.")
            return count

        logger.info(f"Seeding location coordinates (current count={count}, force={force})...")
        if force:
            db.query(LocationCoordinate).delete()
            db.commit()

        airports = load_airports(AIRPORTS_FILE) if os.path.exists(AIRPORTS_FILE) else []
        seaports = load_seaports(SEAPORTS_FILE) if os.path.exists(SEAPORTS_FILE) else []

        all_records = airports + seaports

        # Add supplementary ports if not already present
        existing_codes = {(r["location_type"], r["code"]) for r in all_records}
        for sp in SUPPLEMENTARY_PORTS:
            if (sp["location_type"], sp["code"]) not in existing_codes:
                all_records.append(sp)
                existing_codes.add((sp["location_type"], sp["code"]))

        logger.info(f"Inserting {len(all_records)} total location coordinates into database...")

        # Batch insert in chunks of 1000
        batch_size = 1000
        for i in range(0, len(all_records), batch_size):
            chunk = all_records[i:i + batch_size]
            db.bulk_insert_mappings(LocationCoordinate, chunk)
            db.commit()

        total = db.execute(select(func.count(LocationCoordinate.id))).scalar()
        logger.info(f"Successfully seeded {total} location coordinates!")
        return total
    finally:
        db.close()


if __name__ == "__main__":
    seed_locations_if_needed(force=True)
